"""Загрузка выгрузки, транспорт: сырое тело, кодированное имя, задание и форма отделений.

От пустого корня до страницы, по которой можно считать, — только запросами к странице.
"""

from __future__ import annotations

import http.client
import json
import threading
import time
import urllib.error
import urllib.request
from pathlib import Path
from urllib.parse import quote, urlparse

import openpyxl
import pytest
from conftest import BRANCH_REGIONS, notice_document
from test_page_http import get, post, running

from mkforge.page.server import PageHandler, PageServer, free_port
from mkforge.task.calculate import Places, Waiting


@pytest.fixture
def places(home_in_tmp: Path, distributed_config_path: Path) -> Places:
    root = home_in_tmp
    return Places(
        root=root,
        raw_dir=root / "data" / "raw",
        inputs_dir=root / "data" / "anon",
        config_path=distributed_config_path,
        mapping_path=root / "data" / "contracts.mapping.json",
        work_dir=root / "out" / "рабочие",
        out_dir=root / "out",
    )


@pytest.fixture
def server(places: Places):
    server = PageServer(("127.0.0.1", free_port()), PageHandler, places.open(), places=places)
    for address in running(server):
        yield server, address


def send(address: str, body: bytes, name: str | bytes | None, quoted: bool = True):
    """Файл сырым телом. Имя кодируется, как это делает страница."""
    parsed = urlparse(address)
    connection = http.client.HTTPConnection(parsed.hostname, parsed.port, timeout=30)
    try:
        connection.putrequest("POST", "/api/upload")
        connection.putheader("Content-Type", "application/octet-stream")
        connection.putheader("Content-Length", str(len(body)))
        if name is not None:
            value = quote(name, safe="") if quoted and isinstance(name, str) else name
            connection.putheader("X-File-Name", value)
        connection.endheaders(body)
        response = connection.getresponse()
        return response.status, json.loads(response.read())
    finally:
        connection.close()


def wait_job(address: str, job: str, limit: float = 30) -> dict:
    deadline = time.monotonic() + limit
    while time.monotonic() < deadline:
        _, body, _ = get(address, f"/api/jobs/{job}")
        payload = json.loads(body)
        if payload["state"] != "running":
            return payload
        time.sleep(0.05)
    raise AssertionError("обработка не кончилась")


def refused(call) -> tuple[int, dict]:
    with pytest.raises(urllib.error.HTTPError) as error:
        call()
    return error.value.code, json.loads(error.value.read())


def raw_leftovers(places: Places) -> list[Path]:
    return list(places.raw_dir.iterdir()) if places.raw_dir.is_dir() else []


def test_from_empty_root_to_a_page_that_calculates(server, places, export, tmp_path):
    _, address = server
    _, body, _ = get(address, "/api/state")
    assert json.loads(body)["ready"] is False

    name = "Выгрузка за август; «пул».xlsx"
    status, answer = send(address, export.read_bytes(), name)
    assert status == 200 and answer == {
        "ok": True, "kind": "export", "name": name, "size": export.stat().st_size,
    }, "кириллица и знаки в имени доходят целыми"
    notice = notice_document(tmp_path / "Уведомление.docx")
    status, answer = send(address, notice.read_bytes(), notice.name)
    assert status == 200 and answer["kind"] == "notice"

    status, started = post(address, "/api/prepare", {})
    assert status == 202
    job = wait_job(address, started["job"])
    assert job["state"] == "done", job["error"]
    assert job["name"] == "обработка выгрузки"
    assert job["result"]["ready"] is False, "таблицы отделений еще нет"
    assert raw_leftovers(places) == [], "сырые файлы после обработки не остаются"

    _, body, _ = get(address, "/api/state")
    waiting = json.loads(body)["waiting"]
    assert waiting["missing"] == ["branch_regions.csv"]

    _, body, _ = get(address, "/api/branches")
    described = json.loads(body)
    assert described["missing"] == sorted(BRANCH_REGIONS)

    status, answer = post(address, "/api/branches", {"pairs": dict(BRANCH_REGIONS)})
    assert status == 200 and answer["ok"] is True and answer["ready"] is True

    _, body, _ = get(address, "/api/state")
    state = json.loads(body)
    assert state["ready"] is True and state["answer"]["plan"]["pool_size"] > 0


def test_second_upload_replaces_data_and_the_page_rereads_it(server, places, export, tmp_path):
    _, address = server
    send(address, export.read_bytes(), export.name)
    wait_job(address, post(address, "/api/prepare", {})[1]["job"])
    post(address, "/api/branches", {"pairs": dict(BRANCH_REGIONS)})
    _, body, _ = get(address, "/api/state")
    before = json.loads(body)["answer"]["plan"]["pool_size"]

    wb = openpyxl.load_workbook(export)
    for sheet, header in (("Транзакции участников", 1), ("База участников", 1)):
        ws = wb[sheet]
        contract = ws.cell(2, 1).value
        rows = [row for row in range(header + 1, ws.max_row + 1) if ws.cell(row, 1).value == contract]
        for row in reversed(rows):
            ws.delete_rows(row)
    smaller = tmp_path / "меньше.xlsx"
    wb.save(smaller)

    send(address, smaller.read_bytes(), smaller.name)
    job = wait_job(address, post(address, "/api/prepare", {})[1]["job"])
    assert job["state"] == "done" and job["result"]["ready"] is True
    _, body, _ = get(address, "/api/state")
    assert json.loads(body)["answer"]["plan"]["pool_size"] == before - 1


def test_failed_processing_keeps_the_page_on_old_data(server, places, export, tmp_path, capsys):
    _, address = server
    send(address, export.read_bytes(), export.name)
    wait_job(address, post(address, "/api/prepare", {})[1]["job"])
    post(address, "/api/branches", {"pairs": dict(BRANCH_REGIONS)})
    before = (places.inputs_dir / "transactions.csv").read_bytes()

    broken = tmp_path / "не та.xlsx"
    openpyxl.Workbook().save(broken)
    send(address, broken.read_bytes(), broken.name)
    job = wait_job(address, post(address, "/api/prepare", {})[1]["job"])
    assert job["state"] == "failed" and "нет листа" in job["error"]
    assert (places.inputs_dir / "transactions.csv").read_bytes() == before
    assert raw_leftovers(places) == []
    _, body, _ = get(address, "/api/state")
    assert json.loads(body)["ready"] is True, "по прежним данным считать можно"
    assert f"обработка выгрузки {job['id']} упала: ExportError" in capsys.readouterr().out


def test_unencoded_cyrillic_name_is_refused(server, places, export):
    """http.server читает заголовки в latin-1: кириллица как есть — отказ, а не кракозябры."""
    _, address = server
    status, answer = send(address, export.read_bytes(), "выгрузка.xlsx".encode("utf-8"))
    assert status == 400 and "кодированным" in answer["message"]
    assert raw_leftovers(places) == []


def test_path_in_the_name_is_reduced_to_the_base(server, places, export):
    _, address = server
    status, answer = send(address, export.read_bytes(), "../../../выгрузка.xlsx")
    assert status == 200 and answer["name"] == "выгрузка.xlsx"
    [folder] = raw_leftovers(places)
    assert folder.name.startswith(".upload-")
    assert [path.name for path in folder.iterdir()] == ["выгрузка.xlsx"]


@pytest.mark.parametrize("name,body,code", [
    ("таблица.csv", b"PK\x03\x04", 415),
    ("выгрузка.xlsx", b"not a workbook", 415),
    (None, b"PK\x03\x04", 400),
])
def test_refusals_leave_nothing_and_keep_the_server_alive(server, places, tmp_path, name, body, code):
    _, address = server
    notice = notice_document(tmp_path / "уведомление.docx")
    send(address, notice.read_bytes(), notice.name)
    status, answer = send(address, body, name)
    assert status == code and answer["ok"] is False
    assert raw_leftovers(places) == [], "и принятое до отказа уведомление тоже"
    status, _, _ = get(address, "/api/health")
    assert status == 200


def test_over_the_limit(server, monkeypatch):
    _, address = server
    monkeypatch.setattr("mkforge.task.intake.UPLOAD_LIMIT", 1024)
    status, answer = send(address, b"PK\x03\x04" + b"x" * 4096, "выгрузка.xlsx")
    assert status == 413 and answer["ok"] is False


def test_body_without_length_is_refused(server):
    _, address = server
    parsed = urlparse(address)
    connection = http.client.HTTPConnection(parsed.hostname, parsed.port, timeout=10)
    try:
        connection.request("POST", "/api/upload", body=iter([b"PK\x03\x04"]),
                           headers={"X-File-Name": "a.xlsx", "Transfer-Encoding": "chunked"})
        assert connection.getresponse().status == 411
    finally:
        connection.close()


def test_prepare_without_export_refuses_and_drops_the_notice(server, places, tmp_path):
    _, address = server
    notice = notice_document(tmp_path / "уведомление.docx")
    send(address, notice.read_bytes(), notice.name)
    code, answer = refused(lambda: post(address, "/api/prepare", {}))
    assert code == 409 and "выгрузки .xlsx" in answer["message"]
    assert raw_leftovers(places) == [], "уведомление не лежит до следующей загрузки"


def test_upload_and_branches_wait_for_a_running_job(server, places, export):
    """Уведомление принято до начала сборки, выгрузке отказано: в томе не остается ничего."""
    instance, address = server
    notice = notice_document(places.root / "уведомление.docx")
    send(address, notice.read_bytes(), notice.name)
    release = threading.Event()
    instance.jobs.start(lambda report: release.wait(10) and {"ok": True})
    try:
        status, answer = send(address, export.read_bytes(), export.name)
        assert status == 409 and "сборка книги уже идет" in answer["message"]
        code, answer = refused(lambda: post(address, "/api/branches", {"pairs": {}}))
        assert code == 409
    finally:
        release.set()
    assert raw_leftovers(places) == []


def test_page_without_places_does_not_take_uploads(distributed_config_path, tmp_path, export):
    state = Waiting(inputs_dir=tmp_path / "пусто", config_path=distributed_config_path)
    server = PageServer(("127.0.0.1", free_port()), PageHandler, state, tmp_path / "готовые")
    for address in running(server):
        status, answer = send(address, export.read_bytes(), export.name)
        assert status == 409 and answer["ok"] is False
        _, body, _ = get(address, "/api/state")
        assert json.loads(body)["upload"] is False


def test_leftovers_of_a_killed_page_are_swept_on_start(places, distributed_config_path):
    stale = places.raw_dir / ".upload-0123456789ab"
    stale.mkdir(parents=True)
    (stale / "выгрузка.xlsx").write_bytes(b"PK")
    server = PageServer(("127.0.0.1", free_port()), PageHandler, places.open(), places=places)
    server.server_close()
    assert not stale.exists()
