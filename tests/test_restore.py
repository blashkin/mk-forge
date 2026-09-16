"""Возврат настоящих номеров: подмена не должна ломать формулы."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from mkforge.config import load_config
from mkforge.core.loaders import load_inputs
from mkforge.core.models import ModelError
from mkforge.core.validate import RecalculationError, find_soffice, validate
from mkforge.renderers import participants as pool_sheet
from mkforge.renderers import transactions as tx
from mkforge.renderers.xlsx import build
from mkforge.restore import load_mapping, restore_numbers


@pytest.fixture
def built(anon_dir: Path, config_path: Path, tmp_path: Path) -> Path:
    path = tmp_path / "книга.xlsx"
    build(inputs=load_inputs(anon_dir), config=load_config(config_path), path=path)
    return path


@pytest.fixture
def mapping_path(anon_dir: Path, tmp_path: Path) -> Path:
    """Таблица соответствия, которую написал анонимайзер вместе с anon_dir."""
    path = tmp_path / "mapping.json"
    assert path.exists(), "таблицу соответствия пишет фикстура anon_dir"
    return path


def test_real_numbers_come_back(built, mapping_path, tmp_path, pool):
    out = tmp_path / "с номерами.xlsx"
    result = restore_numbers(
        book=built, mapping_path=mapping_path, out=out, recalculate=False
    )

    assert result.contracts == len(pool)
    assert result.replaced[pool_sheet.SHEET] == len(pool)
    assert result.replaced[tx.SHEET] == len(pool) * 2  # по две строки на договор

    workbook = openpyxl.load_workbook(out)
    try:
        written = {
            workbook[pool_sheet.SHEET][f"{pool_sheet.CONTRACT}{row}"].value
            for row in range(2, 2 + len(pool))
        }
        assert written == set(pool)
        assert not any(str(value).startswith("Д-") for value in written)
    finally:
        workbook.close()


def test_transactions_and_pool_stay_consistent(built, mapping_path, tmp_path):
    """SUMIFS сопоставляет две колонки договоров: разъехаться они не должны."""
    out = tmp_path / "с номерами.xlsx"
    restore_numbers(book=built, mapping_path=mapping_path, out=out, recalculate=False)

    workbook = openpyxl.load_workbook(out)
    try:
        pool_numbers = {
            workbook[pool_sheet.SHEET][f"A{row}"].value
            for row in range(2, workbook[pool_sheet.SHEET].max_row + 1)
            if workbook[pool_sheet.SHEET][f"A{row}"].value
        }
        tx_numbers = {
            workbook[tx.SHEET][f"A{row}"].value
            for row in range(2, workbook[tx.SHEET].max_row + 1)
            if workbook[tx.SHEET][f"A{row}"].value
        }
        assert tx_numbers == pool_numbers
    finally:
        workbook.close()


def test_formulas_survive(built, mapping_path, tmp_path):
    out = tmp_path / "с номерами.xlsx"
    restore_numbers(book=built, mapping_path=mapping_path, out=out, recalculate=False)

    workbook = openpyxl.load_workbook(out)
    try:
        for column in ("I", "K", "O", "R", "S", "U"):
            value = workbook[pool_sheet.SHEET][f"{column}2"].value
            assert isinstance(value, str) and value.startswith("=")
        assert workbook["Расчет акции"]["B77"].value.startswith("=IF(")
    finally:
        workbook.close()


def test_unknown_pseudonym_is_reported(built, mapping_path, tmp_path):
    """Книга, собранная на других данных, не должна пройти молча."""
    mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    mapping.pop(next(iter(mapping)))
    mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(ModelError, match="таблице соответствия"):
        restore_numbers(book=built, mapping_path=mapping_path, out=tmp_path / "нет.xlsx")


def test_missing_mapping_names_the_command(built, tmp_path):
    with pytest.raises(ModelError, match="mk-forge prepare"):
        restore_numbers(
            book=built, mapping_path=tmp_path / "нет.json", out=tmp_path / "нет.xlsx"
        )


def test_empty_mapping_is_rejected(built, tmp_path):
    empty = tmp_path / "пусто.json"
    empty.write_text("{}", encoding="utf-8")
    with pytest.raises(ModelError, match="пустая"):
        load_mapping(empty)


def test_book_without_expected_sheet(mapping_path, tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Не то"
    wrong = tmp_path / "не то.xlsx"
    workbook.save(wrong)
    with pytest.raises(ModelError, match="Транзакции участников"):
        restore_numbers(book=wrong, mapping_path=mapping_path, out=tmp_path / "нет.xlsx")


@pytest.mark.libreoffice
def test_restored_book_still_computes(
    anon_dir: Path, config_path: Path, built: Path, mapping_path: Path, tmp_path: Path
):
    """Главное: после подмены номеров книга должна считать те же числа."""
    try:
        find_soffice()
    except RecalculationError:
        pytest.skip("LibreOffice не установлен")

    out = tmp_path / "с номерами.xlsx"
    restore_numbers(book=built, mapping_path=mapping_path, out=out, recalculate=False)

    inputs = load_inputs(anon_dir)
    config = load_config(config_path)
    before = validate(book=built, inputs=inputs, config=config, work_dir=tmp_path / "до")
    assert before.ok, before.report()

    # У книги с настоящими номерами сверять пул по псевдонимам нельзя,
    # поэтому проверяем инвариант по уровням — он от номеров не зависит.
    from mkforge.core.validate import recalculate

    computed = recalculate(out, tmp_path / "после")
    workbook = openpyxl.load_workbook(computed, data_only=True)
    try:
        for index in range(len(config.targets)):
            row = 77 + index
            target = workbook["Расчет акции"][f"AQ{row}"].value
            actual = workbook["Расчет акции"][f"AR{row}"].value
            assert actual == pytest.approx(target, abs=1e-9)
    finally:
        workbook.close()


@pytest.mark.libreoffice
def test_restored_book_carries_numbers(built, mapping_path, tmp_path):
    """Иначе Excel откроет книгу верно, а предпросмотр покажет пустые ячейки."""
    try:
        find_soffice()
    except RecalculationError:
        pytest.skip("LibreOffice не установлен")

    out = tmp_path / "готовая.xlsx"
    result = restore_numbers(book=built, mapping_path=mapping_path, out=out)
    assert result.recalculated, result.report()

    values = openpyxl.load_workbook(out, data_only=True)
    formulas = openpyxl.load_workbook(out)
    try:
        assert isinstance(values[pool_sheet.SHEET]["O2"].value, (int, float))
        assert str(formulas[pool_sheet.SHEET]["O2"].value).startswith("=")
    finally:
        values.close()
        formulas.close()


def test_recalculation_can_be_skipped(built, mapping_path, tmp_path):
    out = tmp_path / "без пересчета.xlsx"
    result = restore_numbers(
        book=built, mapping_path=mapping_path, out=out, recalculate=False
    )
    assert not result.recalculated
    assert "предпросмотр" in result.report()


def test_deliverable_goes_one_level_up(tmp_path):
    """Рабочая копия собирается в подпапке, готовая ложится в out/ под тем же именем."""
    from mkforge.restore import deliverable_path

    out = tmp_path / "out"
    book = out / "рабочие" / "Книга.xlsx"
    assert deliverable_path(book, out) == out / "Книга.xlsx"


def test_deliverable_never_overwrites_its_source(tmp_path):
    """Если книга уже в папке готовых, имя дополняется: иначе затрем исходник."""
    from mkforge.restore import deliverable_path

    out = tmp_path / "out"
    out.mkdir()
    book = out / "Книга.xlsx"
    book.write_text("", encoding="utf-8")
    target = deliverable_path(book, out)
    assert target != book
    assert target.name == "Книга с номерами.xlsx"


def test_ready_book_appears_in_one_rename(built, mapping_path, tmp_path, monkeypatch):
    """Пока книга пересчитывается, в готовых ее нет: прерванная сборка не оставит
    там недосчитанную книгу, которую страница предложила бы скачать."""
    ready = tmp_path / "готовые"
    out = ready / "Книга.xlsx"

    def recalculate(book, out_dir, profile_dir=None):
        assert not out.exists(), "книга легла в готовые до пересчета"
        assert book.parent.parent == ready, "черновик рядом с целью: переименование в пределах диска"
        out_dir.mkdir(parents=True, exist_ok=True)
        computed = out_dir / book.name
        computed.write_bytes(book.read_bytes())
        return computed

    monkeypatch.setattr("mkforge.core.validate.recalculate", recalculate)
    result = restore_numbers(book=built, mapping_path=mapping_path, out=out)

    assert result.recalculated and result.path == out
    assert [path.name for path in ready.iterdir()] == ["Книга.xlsx"], "временная папка убрана"


def test_interrupted_recalculation_leaves_nothing_ready(built, mapping_path, tmp_path, monkeypatch):
    ready = tmp_path / "готовые"

    def recalculate(book, out_dir, profile_dir=None):
        raise KeyboardInterrupt  # страницу остановили посреди пересчета

    monkeypatch.setattr("mkforge.core.validate.recalculate", recalculate)
    with pytest.raises(KeyboardInterrupt):
        restore_numbers(book=built, mapping_path=mapping_path, out=ready / "Книга.xlsx")
    assert not list(ready.iterdir())
