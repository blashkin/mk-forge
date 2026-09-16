"""Пересчет собранной книги и проверка ее инвариантов.

Формулы в книге и расчетное ядро считают одно и то же двумя путями. Ядро уже
сверено с эталоном, поэтому здесь проверяется третье: что формулы, которые
записал рендерер, дают те же числа, что ядро.

Пересчет делает LibreOffice в режиме без окна. Свежесобранная книга не содержит
ни одного закешированного значения, поэтому LibreOffice обязан вычислить их все —
настройка «пересчитывать при открытии» на это не влияет.
"""

from __future__ import annotations

import os
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from mkforge.campaigns.motivational import (
    Assumptions,
    build_plan,
    forecast_pool_from,
    plan_from_config,
)
from mkforge.config import CampaignConfig
from mkforge.core.loaders import load_inputs
from mkforge.core.models import Inputs, ModelError
from mkforge.core.pool import build_pool
from mkforge.core.verify import Comparison, _relative
from mkforge.renderers import calculation as calc
from mkforge.renderers import distribution as dist
from mkforge.renderers import participants as pool_sheet

SOFFICE_ENV = "MK_FORGE_SOFFICE"
RECALC_TIMEOUT = 600
DEFAULT_TOLERANCE = 1e-9

# Где искать LibreOffice, если его нет на PATH. Linux нужен для образа,
# Windows — потому что на нем установщик не добавляет soffice в PATH.
SOFFICE_CANDIDATES = (
    # macOS
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    # Linux
    "/usr/bin/soffice",
    "/usr/lib/libreoffice/program/soffice",
    "/opt/libreoffice/program/soffice",
    # Windows
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
)

ERROR_MARKERS = ("#DIV/0!", "#VALUE!", "#REF!", "#N/A", "#NAME?", "#NULL!", "#NUM!")

# Колонка «Базы участников» -> поле участника пула.
POOL_COLUMNS = {
    "среднемес. объем топлива": (pool_sheet.FUEL_LITERS, "fuel_liters_per_month"),
    "СТП продукта": (pool_sheet.STP, "stp_rate"),
    "среднемес. объем продукта": (pool_sheet.PRODUCT_TONS, "product_tons_per_month"),
    "ставка сервисного сбора": (pool_sheet.FEE_RATE, "service_fee_rate"),
}

# Колонка раздела 5 -> как ту же величину зовет ядро.
LEVEL_COLUMNS = {
    "надбавка": ("B", lambda level: level.markup),
    "объем доп. эффекта": ("F", lambda level: level.effect.tons),
    "затраты акции": ("G", lambda level: level.effect.costs),
    "сервисный сбор доп.": ("H", lambda level: level.effect.service_fee),
    "маржа доп.": ("L", lambda level: level.effect.margin),
    "окупаемость": ("M", lambda level: level.effect.payback),
    "ROI": ("N", lambda level: level.effect.roi),
    "эффективная скидка": ("AR", lambda level: level.effective_discount),
    "итоговая скидка": ("AS", lambda level: level.with_campaign.total_discount_rate),
}


# Колонка раздела 6 -> как ту же величину зовет ядро. Срез и итог считаются
# одинаково, поэтому таблица одна на оба случая.
DEPTH_COLUMNS = {
    "надбавка среза": (dist.MARKUP, lambda part: part.markup),
    "затраты среза": (dist.EXTRA_COSTS, lambda part: part.effect.costs),
    "маржа среза": (dist.EXTRA_MARGIN, lambda part: part.effect.margin),
    "окупаемость среза": (dist.PAYBACK, lambda part: part.effect.payback),
    "ROI среза": (dist.ROI, lambda part: part.effect.roi),
    "эффективная скидка среза": (
        dist.EFFECTIVE, lambda part: part.with_campaign.effective_discount
    ),
}


class RecalculationError(Exception):
    """Пересчитать книгу не удалось."""


@dataclass
class ValidationReport:
    """Итог проверки собранной книги."""

    book: Path
    tolerance: float = DEFAULT_TOLERANCE
    formula_errors: dict[str, int] = field(default_factory=dict)
    empty_cells: list[str] = field(default_factory=list)
    levels_checked: int = 0
    depths_checked: int = 0
    target_equals_actual: bool = True
    depth_equals_actual: bool = True
    breakeven: str = ""
    comparisons: list[Comparison] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        if self.formula_errors or self.empty_cells or not self.target_equals_actual:
            return False
        if not self.depth_equals_actual:
            return False
        return all(c.matches(self.tolerance) for c in self.comparisons)

    def report(self) -> str:
        lines = [f"книга: {self.book}"]
        if self.formula_errors:
            lines += [f"  ОШИБКА В ФОРМУЛАХ {kind}: {count}"
                      for kind, count in sorted(self.formula_errors.items())]
        else:
            lines.append("  ошибок в формулах нет")
        if self.empty_cells:
            lines.append(f"  НЕ ПОСЧИТАЛИСЬ ячейки: {', '.join(self.empty_cells[:6])}")
        lines.append(
            f"  уровней сверено: {self.levels_checked}, "
            + ("цель равна фактической эффективной скидке" if self.target_equals_actual
               else "ЦЕЛЬ НЕ РАВНА ФАКТУ")
        )
        if self.depths_checked:
            lines.append(
                f"  срезов распределения сверено: {self.depths_checked}, "
                + ("заданная глубина равна фактической" if self.depth_equals_actual
                   else "ГЛУБИНА СРЕЗА НЕ РАВНА ФАКТУ")
            )
        if self.breakeven:
            lines.append(f"  {self.breakeven}")
        for comparison in self.comparisons:
            if comparison.matches(self.tolerance):
                lines.append(f"  {comparison.name}: совпадает ({comparison.compared} значений)")
            else:
                lines.append(
                    f"  {comparison.name}: РАСХОЖДЕНИЕ до {comparison.worst_diff:.4%}"
                    f" на {comparison.worst_contract}"
                )
        lines.append("итог: книга сходится" if self.ok else "итог: книгу отдавать нельзя")
        return "\n".join(lines)


def find_soffice() -> Path:
    """Найти LibreOffice. Путь можно задать переменной MK_FORGE_SOFFICE."""
    if override := os.environ.get(SOFFICE_ENV):
        path = Path(override)
        if not path.exists():
            raise RecalculationError(f"{SOFFICE_ENV} указывает на {path}, которого нет")
        return path
    for name in ("soffice", "soffice.exe"):
        if found := shutil.which(name):
            return Path(found)
    for candidate in SOFFICE_CANDIDATES:
        if Path(candidate).exists():
            return Path(candidate)
    raise RecalculationError(
        "LibreOffice не найден. Поставь его (brew install --cask libreoffice) "
        f"или укажи путь в {SOFFICE_ENV}"
    )


def recalc_command(soffice: Path, book: Path, out_dir: Path, profile: Path) -> list[str]:
    """Команда пересчета. Вынесена отдельно, чтобы ее проверял тест.

    Профиль задается URL, и он обязан быть абсолютным и закодированным: от
    относительного пути LibreOffice принимает первый сегмент за имя хоста
    и молча зависает, а кириллица в пути ломает URL без процентного кодирования.
    """
    return [
        str(soffice),
        "--headless",
        "--norestore",
        # Отдельный профиль: иначе пересчет не запустится, если LibreOffice уже открыт.
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to",
        "xlsx:Calc MS Excel 2007 XML",
        "--outdir",
        str(out_dir.resolve()),
        str(book.resolve()),
    ]


def recalculate(book: Path, out_dir: Path, profile_dir: Path | None = None) -> Path:
    """Пересчитать книгу и вернуть путь к файлу со значениями."""
    soffice = find_soffice()
    book = book.resolve()
    out_dir = out_dir.resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    # Профиль задается URL, и он обязан быть абсолютным и закодированным:
    # от относительного пути LibreOffice принимает первый сегмент за имя хоста
    # и молча зависает, а кириллица в пути ломает URL без процентного кодирования.
    profile = (profile_dir or out_dir / "profile").resolve()
    profile.mkdir(parents=True, exist_ok=True)

    command = recalc_command(soffice, book, out_dir, profile)
    try:
        result = subprocess.run(
            command, capture_output=True, text=True, timeout=RECALC_TIMEOUT, check=False
        )
    except subprocess.TimeoutExpired as error:
        raise RecalculationError(f"пересчет не уложился в {RECALC_TIMEOUT} с") from error

    recalculated = out_dir / book.name
    if not recalculated.exists():
        raise RecalculationError(
            f"LibreOffice не создал файл. Код {result.returncode}. "
            f"{(result.stderr or result.stdout or '').strip()[:400]}"
        )
    return recalculated


def _scan_errors(workbook) -> dict[str, int]:
    """Сколько ячеек с каждой ошибкой формулы во всей книге."""
    errors: dict[str, int] = {}
    for sheet in workbook.sheetnames:
        for row in workbook[sheet].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERROR_MARKERS:
                    errors[cell.value] = errors.get(cell.value, 0) + 1
    return errors


def validate(
    book: Path,
    inputs: Inputs,
    config: CampaignConfig,
    work_dir: Path,
    tolerance: float = DEFAULT_TOLERANCE,
) -> ValidationReport:
    """Пересчитать книгу и сверить ее числа с расчетным ядром."""
    recalculated = recalculate(book, work_dir)
    report = ValidationReport(book=book, tolerance=tolerance)

    # Если задано распределение глубины, прогноз новых приходит из прогнозного
    # пула строками, и таблица уровней в книге считает по нему же. Сверять ее
    # с прогнозом из плана было бы сверкой двух разных расчетов.
    distributed = plan_from_config(inputs, config) if config.distribution else None
    plan = build_plan(
        inputs=inputs,
        product=config.product,
        start=config.start,
        end=config.end,
        assumptions=Assumptions(
            plan_participants=config.plan_participants,
            share_without=config.share_without,
            share_with=config.share_with,
            comms_per_client=config.comms_per_client,
            new_share_without=config.new_share_without,
            large_threshold_thousand_liters=config.large_threshold_thousand_liters,
            large_shift_points=config.large_shift_points,
        ),
        targets=config.targets,
        new=(
            forecast_pool_from(inputs, config).new_clients()
            if config.distribution
            else None
        ),
    )
    members = {m.contract: m for m in build_pool(inputs, config.product).members}

    workbook = openpyxl.load_workbook(recalculated, data_only=True)
    try:
        report.formula_errors = _scan_errors(workbook)

        # Раздел 5: уровни глубины скидки.
        ws = workbook[calc.SHEET]
        level_comparisons = {name: Comparison(name=name) for name in LEVEL_COLUMNS}
        for index, level in enumerate(plan.levels):
            row = calc.LEVELS_FIRST_ROW + index
            report.levels_checked += 1
            target = ws[f"AQ{row}"].value
            actual = ws[f"AR{row}"].value
            if not isinstance(target, (int, float)) or not isinstance(actual, (int, float)):
                report.target_equals_actual = False
            elif abs(target - actual) > tolerance:
                report.target_equals_actual = False
            for name, (column, extract) in LEVEL_COLUMNS.items():
                value = ws[f"{column}{row}"].value
                if not isinstance(value, (int, float)):
                    report.empty_cells.append(f"{calc.SHEET}!{column}{row}")
                    continue
                comparison = level_comparisons[name]
                comparison.compared += 1
                diff = _relative(extract(level), value)
                if diff > comparison.worst_diff:
                    comparison.worst_diff = diff
                    comparison.worst_contract = f"уровень {level.target * 100:g}%"
        report.comparisons.extend(level_comparisons.values())

        # Раздел 6: распределение глубины скидки и его итог.
        if distributed is not None:
            _check_distribution(ws, distributed, report, tolerance)

        # «База участников»: по договору на строку.
        ws = workbook[pool_sheet.SHEET]
        pool_comparisons = {name: Comparison(name=name) for name in POOL_COLUMNS}
        row = pool_sheet.FIRST_ROW
        while True:
            contract = ws[f"{pool_sheet.CONTRACT}{row}"].value
            if contract in (None, ""):
                break
            member = members.get(str(contract))
            if member is None:
                raise ModelError(f"в книге договор {contract}, которого нет в данных")
            for name, (column, attribute) in POOL_COLUMNS.items():
                value = ws[f"{column}{row}"].value
                if not isinstance(value, (int, float)):
                    report.empty_cells.append(f"{pool_sheet.SHEET}!{column}{row}")
                    continue
                comparison = pool_comparisons[name]
                comparison.compared += 1
                diff = _relative(getattr(member, attribute), value)
                if diff > comparison.worst_diff:
                    comparison.worst_diff = diff
                    comparison.worst_contract = str(contract)
            row += 1
        report.comparisons.extend(pool_comparisons.values())
    finally:
        workbook.close()

    return report


def _check_distribution(ws, plan, report: "ValidationReport", tolerance: float) -> None:
    """Сверить раздел распределения: каждый срез и строку «Итого».

    Итог сверяется отдельно от срезов, а не как их сумма: сумма сошлась бы
    и при взаимно компенсирующих ошибках в срезах.
    """
    count = len(plan.depths)
    comparisons = {name: Comparison(name=name) for name in DEPTH_COLUMNS}

    parts = [
        (f"глубина {part.depth * 100:g}%", dist.FIRST_ROW + index, part)
        for index, part in enumerate(plan.depths)
    ]
    parts.append(("итого по акции", dist.total_row(count), plan))

    for label, row, part in parts:
        report.depths_checked += 1
        expected_depth = getattr(part, "depth", None)
        actual = ws[f"{dist.EFFECTIVE}{row}"].value
        if expected_depth is not None:
            if not isinstance(actual, (int, float)) or abs(actual - expected_depth) > tolerance:
                report.depth_equals_actual = False
        for name, (column, extract) in DEPTH_COLUMNS.items():
            value = ws[f"{column}{row}"].value
            if not isinstance(value, (int, float)):
                report.empty_cells.append(f"{calc.SHEET}!{column}{row}")
                continue
            comparison = comparisons[name]
            comparison.compared += 1
            diff = _relative(extract(part), value)
            if diff > comparison.worst_diff:
                comparison.worst_diff = diff
                comparison.worst_contract = label
    report.comparisons.extend(comparisons.values())

    # Порог безубыточности: книга считает его формулой, ядро — решением
    # того же уравнения. Расходиться им нельзя.
    rows = list(dist.breakeven_rows(count).values())
    for row, expected, name in (
        (rows[4], plan.breakeven_payback, "порог окупаемости"),
        (rows[5], plan.breakeven_roi, "порог ROI"),
    ):
        value = ws[f"B{row}"].value
        if expected is None:
            if isinstance(value, (int, float)):
                report.breakeven = (
                    f"РАСХОЖДЕНИЕ: {name} в книге {value}, ядро считает его недостижимым"
                )
            continue
        if not isinstance(value, (int, float)) or abs(value - expected) > 1e-6:
            report.breakeven = (
                f"РАСХОЖДЕНИЕ: {name} в книге {value}, в ядре {expected}"
            )
            report.depth_equals_actual = False
    if not report.breakeven:
        report.breakeven = "порог безубыточности в книге совпадает с ядром"


def validate_path(
    book: Path,
    inputs_dir: Path,
    config_path: Path,
    work_dir: Path,
    tolerance: float = DEFAULT_TOLERANCE,
) -> ValidationReport:
    """Обертка для командной строки: читает данные и конфиг сама."""
    from mkforge.config import load_config

    return validate(
        book=book,
        inputs=load_inputs(inputs_dir),
        config=load_config(config_path),
        work_dir=work_dir,
        tolerance=tolerance,
    )
