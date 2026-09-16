"""Пересчет правки и поиск ошибок в формулах.

Пересчитывает книгу в LibreOffice и печатает, сколько ячеек с каждой ошибкой
формулы. Адреса ячеек с ошибками печатаются, содержимое — нет.

Код возврата: 0 — ошибок нет, 1 — есть.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl

from mkforge.core.validate import ERROR_MARKERS, recalculate


def errors(book: Path) -> dict[str, list[str]]:
    """Адреса ячеек с ошибками формул, по виду ошибки."""
    workbook = openpyxl.load_workbook(book, data_only=True)
    try:
        found: dict[str, list[str]] = {}
        for title in workbook.sheetnames:
            for row in workbook[title].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in ERROR_MARKERS:
                        found.setdefault(cell.value, []).append(f"{title}!{cell.coordinate}")
        return found
    finally:
        workbook.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Пересчитать правку и найти ошибки формул")
    parser.add_argument("правка", type=Path, help="книга после правки")
    parser.add_argument("--куда", type=Path, default=Path("out/правка/пересчет"),
                        dest="out_dir", help="папка для пересчитанной копии")
    parser.add_argument("--предел", type=int, default=20, dest="limit",
                        help="сколько адресов печатать на вид ошибки")
    args = parser.parse_args(argv)

    recalculated = recalculate(args.правка, args.out_dir)
    found = errors(recalculated)
    if not found:
        print(f"Пересчитано: {recalculated}")
        print("Ошибок в формулах нет.")
        return 0

    print(f"Пересчитано: {recalculated}")
    for marker, cells in sorted(found.items()):
        print(f"{marker}: {len(cells)} — {', '.join(cells[:args.limit])}"
              + (" ..." if len(cells) > args.limit else ""))
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
