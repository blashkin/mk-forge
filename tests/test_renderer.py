"""Сборка книги: адреса как в эталоне, диапазоны от данных, значений нет."""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pytest

from mkforge.config import load_config
from mkforge.core.loaders import TRANSACTIONS_FILE, load_inputs
from mkforge.renderers import participants, parameters, transactions
from mkforge.renderers.xlsx import SHEET_ORDER, build


@pytest.fixture
def book(anon_dir: Path, config_path: Path, tmp_path: Path) -> Path:
    path = tmp_path / "книга.xlsx"
    build(inputs=load_inputs(anon_dir), config=load_config(config_path), path=path)
    return path


@pytest.fixture
def sheets(book: Path):
    workbook = openpyxl.load_workbook(book)
    yield workbook
    workbook.close()


def test_all_sheets_in_reference_order(sheets, config_path: Path):
    """Лист прогнозного пула появляется только при заданном распределении."""
    from mkforge.renderers.xlsx import sheet_order

    config = load_config(config_path)
    expected = sheet_order(config.distribution is not None)
    assert sheets.sheetnames == list(expected)
    assert expected[0] == "Витрина акции", (
        "без распределения первой остается витрина, как в эталоне"
    )


def test_transactions_are_laid_out(sheets, anon_dir):
    ws = sheets[transactions.SHEET]
    count = len(load_inputs(anon_dir).transactions)
    assert ws["A1"].value == "№ договора"
    assert ws["W1"].value == "Отделение ТО"
    assert ws[f"A{transactions.last_row(count)}"].value is not None
    assert ws[f"A{transactions.last_row(count) + 1}"].value is None


def test_ranges_are_computed_from_row_count(sheets, anon_dir):
    """Дефект эталона не воспроизведен: все диапазоны кончаются на последней строке."""
    count = len(load_inputs(anon_dir).transactions)
    expected = transactions.last_row(count)

    ws = sheets[participants.SHEET]
    bounds = set()
    for column in ("I", "O", "P", "R"):
        formula = ws[f"{column}2"].value
        bounds.update(
            int(bound)
            for bound in re.findall(r"Транзакции участников'!\$[A-Z]+\$2:\$[A-Z]+\$(\d+)", formula)
        )
    assert bounds == {expected}


def test_more_transactions_move_the_bounds(anon_dir, config_path, tmp_path):
    """Добавилась строка — сдвинулась граница. Именно этого эталон не умел."""
    path = anon_dir / TRANSACTIONS_FILE
    lines = path.read_text(encoding="utf-8").splitlines()
    lines.append(lines[1])
    path.write_text("\n".join(lines), encoding="utf-8")

    out = tmp_path / "больше.xlsx"
    build(inputs=load_inputs(anon_dir), config=load_config(config_path), path=out)
    workbook = openpyxl.load_workbook(out)
    formula = workbook[participants.SHEET]["O2"].value
    workbook.close()
    assert f"$L$2:$L${transactions.last_row(len(lines) - 1)}" in formula


def test_percentages_are_stored_as_fractions(sheets, anon_dir):
    """Шкала в долях, а не числами: скрытого деления на 100 в формулах нет."""
    ws = sheets[parameters.SHEET]
    scale = load_inputs(anon_dir).scale
    for offset, bracket in enumerate(scale.brackets):
        assert ws[f"C{parameters.SCALE_FIRST_ROW + offset}"].value == pytest.approx(bracket.dt)
    # Ищем деление ровно на 100, не задевая перевод литров в тысячи (/1000).
    lookup = sheets[participants.SHEET]["K2"].value
    assert not re.search(r"/100(?!\d)", lookup)
    assert "/1000" in lookup  # тысячи литров для поиска сегмента — это законно


def test_branch_dictionary_is_written(sheets, anon_dir):
    """Карта отделений живет отдельным блоком, а не внутри формулы."""
    ws = sheets[parameters.SHEET]
    expected = load_inputs(anon_dir).branch_regions
    written = {}
    for offset in range(len(expected)):
        row = parameters.BRANCHES_FIRST_ROW + offset
        written[ws[f"Q{row}"].value] = ws[f"R{row}"].value
    assert written == expected


def test_monthly_margin_is_readable(sheets):
    """Формула маржи месяца должна быть обозримой, а не на 2600 символов."""
    formula = sheets[parameters.SHEET]["E36"].value
    assert formula.startswith("=IF(COUNTIFS(")
    assert len(formula) < 400
    assert "SUMPRODUCT" in formula


def test_no_cached_values(book):
    """Книга содержит формулы, а не результаты: считать должен Excel."""
    workbook = openpyxl.load_workbook(book, data_only=True)
    try:
        assert workbook[participants.SHEET]["O2"].value is None
        assert workbook[parameters.SHEET]["C29"].value is None
    finally:
        workbook.close()


def test_computed_cells_are_formulas_not_numbers(sheets):
    """Ни одного захардкоженного результата: считаемое записано формулой."""
    ws = sheets[participants.SHEET]
    for column in ("D", "I", "K", "L", "M", "N", "O", "R", "S", "T", "U", "V"):
        value = ws[f"{column}2"].value
        assert isinstance(value, str) and value.startswith("="), f"{column}2 не формула"


def test_input_values_are_written_as_numbers(sheets, anon_dir):
    """Входное — числами: цена и OPEX приходят из данных, а не считаются."""
    ws = sheets[parameters.SHEET]
    economics = load_inputs(anon_dir).economics["ДТ"]
    assert ws["C26"].value == pytest.approx(economics.gross_revenue)
    assert ws["C28"].value == pytest.approx(economics.opex)


def test_participant_rows_match_pool(sheets, pool):
    ws = sheets[participants.SHEET]
    assert ws[f"A{participants.last_row(len(pool))}"].value is not None
    assert ws[f"A{participants.last_row(len(pool)) + 1}"].value is None
