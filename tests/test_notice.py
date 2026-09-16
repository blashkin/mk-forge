"""Разбор уведомления о СТП: нормализация чисел и границ, поиск таблицы."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from mkforge.core.notice import (
    OPEN_BOUND,
    NoticeError,
    parse_bounds,
    parse_notice,
    parse_rate,
)

# Как записано в уведомлении -> что должно получиться.
RATES = {
    "-3,50": 0.035,
    "-4,50": 0.045,
    "0,00": 0.0,
    "-1,00": 0.01,
    "": 0.0,
}
BOUNDS = {
    "0 – 5**": (0, 5),  # длинное тире и сноска
    "5 - 10": (5, 10),
    "700 - 1 000": (700, 1000),  # пробел как разделитель тысяч
    "1 000 - 2 000": (1000, 2000),
    "более 3 000": (3000, OPEN_BOUND),
    "свыше 500": (500, OPEN_BOUND),
    "до 5": (0, 5),
}


@pytest.mark.parametrize("text,expected", RATES.items())
def test_rate_becomes_positive_fraction(text, expected):
    """В уведомлении скидка — отрицательная поправка, в расчете — положительная доля."""
    assert parse_rate(text) == pytest.approx(expected)


def test_unparsable_rate(text="три с половиной"):
    with pytest.raises(NoticeError, match="не разобрать скидку"):
        parse_rate(text)


@pytest.mark.parametrize("text,expected", BOUNDS.items())
def test_bounds_are_normalized(text, expected):
    assert parse_bounds(text) == expected


def test_bounds_without_digits():
    with pytest.raises(NoticeError, match="границы объема"):
        parse_bounds("по договоренности")


def notice_document(path: Path, *, highway: bool = True, rows=None) -> Path:
    """Синтетическое уведомление: таблица-пустышка и таблица шкалы, как в настоящем."""
    document = docx.Document()

    decoy = document.add_table(rows=2, cols=2)
    decoy.rows[0].cells[0].text = "Руководителю"

    products = ["АБ", "СУГ", "ДТ"] + (["ДТ на трассовых и автоматических АЗС"] if highway else [])
    table = document.add_table(rows=2, cols=2 + len(products))
    table.rows[0].cells[0].text = "Торговая Точка"
    table.rows[0].cells[1].text = "Объем выборки НП (АБ, СУГ, ДТ) клиента"
    table.rows[1].cells[0].text = "Торговая Точка"
    table.rows[1].cells[1].text = "Объем выборки НП"
    for offset, product in enumerate(products):
        table.rows[1].cells[2 + offset].text = product

    for bounds, rates in rows or [("0 – 5**", ("0,00", "0,00", "-3,50", "-4,50")),
                                   ("5 - 10", ("0,00", "0,00", "-3,50", "-4,50")),
                                   ("более 10", ("0,00", "0,00", "-1,00", "-2,00"))]:
        row = table.add_row()
        row.cells[0].text = "АЗС, РФ"
        row.cells[1].text = bounds
        for offset in range(len(products)):
            row.cells[2 + offset].text = rates[offset]

    document.save(str(path))
    return path


def test_parses_the_scale(tmp_path):
    brackets = parse_notice(notice_document(tmp_path / "уведомление.docx"))
    assert len(brackets) == 3
    assert [b.low for b in brackets] == [0, 5, 10]
    assert brackets[-1].high == OPEN_BOUND
    assert brackets[0].rates["дт"] == pytest.approx(0.035)
    assert brackets[0].rates["дт_трасса"] == pytest.approx(0.045)
    assert brackets[0].rates["аб"] == 0.0


def test_footnote_markers_are_stripped_from_label(tmp_path):
    brackets = parse_notice(notice_document(tmp_path / "уведомление.docx"))
    assert "*" not in brackets[0].label


def test_decoy_table_is_skipped(tmp_path):
    """В документе есть таблица с адресатами — она не должна сойти за шкалу."""
    brackets = parse_notice(notice_document(tmp_path / "уведомление.docx"))
    assert all(bracket.rates for bracket in brackets)


def test_highway_column_is_optional(tmp_path):
    rows = [("0 – 5", ("0,00", "0,00", "-3,50")), ("более 5", ("0,00", "0,00", "-1,00"))]
    brackets = parse_notice(
        notice_document(tmp_path / "без трассы.docx", highway=False, rows=rows)
    )
    assert "дт_трасса" not in brackets[0].rates
    assert brackets[0].rates["дт"] == pytest.approx(0.035)


def test_row_for_csv_is_flat(tmp_path):
    row = parse_notice(notice_document(tmp_path / "уведомление.docx"))[0].row()
    assert set(row) >= {"сегмент", "мин_тыс_л", "макс_тыс_л", "аб", "суг", "дт", "дт_трасса"}
    assert not any(isinstance(value, dict) for value in row.values())


def test_document_without_scale_table(tmp_path):
    document = docx.Document()
    document.add_table(rows=2, cols=2).rows[0].cells[0].text = "Руководителю"
    path = tmp_path / "не уведомление.docx"
    document.save(str(path))
    with pytest.raises(NoticeError, match="не нашлась таблица шкалы"):
        parse_notice(path)


def test_missing_file(tmp_path):
    with pytest.raises(NoticeError, match="нет уведомления"):
        parse_notice(tmp_path / "нет-такого.docx")
