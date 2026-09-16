"""Сверка расчетного ядра с эталонной книгой по каждому договору.

Ядро и формулы считают одно и то же двумя путями. Эта сверка ловит расхождение
до того, как оно попадет в книгу для экономистов.

Работает только локально: читает книгу с настоящими номерами договоров и переводит
их в псевдонимы через таблицу соответствия. Ни книга, ни таблица за пределы машины
не уходят.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from mkforge.campaigns.motivational import Assumptions, build_plan
from mkforge.config import CampaignConfig
from mkforge.core.anonymize import CONTRACTS_SHEET
from mkforge.core.loaders import load_inputs
from mkforge.core.models import Inputs, ModelError
from mkforge.core.pool import Pool, build_pool

CALCULATION_SHEET = "Расчет акции"
LAST_ROW = 1453
LEVEL_FIRST_ROW = 77
DEFAULT_TOLERANCE = 1e-9

# Колонка «Базы участников» -> поле участника пула.
POOL_COLUMNS = {
    "среднемес. объем топлива": ("I", "fuel_liters_per_month"),
    "СТП продукта": ("K", "stp_rate"),
    "среднемес. объем продукта": ("O", "product_tons_per_month"),
    "ставка сервисного сбора": ("R", "service_fee_rate"),
}


# Колонка раздела уровней «Расчета акции» -> как ту же величину зовет ядро.
LEVEL_COLUMNS = {
    "надбавка": ("B", lambda level: level.markup),
    "объем доп. эффекта": ("F", lambda level: level.effect.tons),
    "затраты акции": ("G", lambda level: level.effect.costs),
    "сервисный сбор доп.": ("H", lambda level: level.effect.service_fee),
    "маржа доп.": ("L", lambda level: level.effect.margin),
    "окупаемость": ("M", lambda level: level.effect.payback),
    "ROI": ("N", lambda level: level.effect.roi),
    "эффективная скидка": ("AR", lambda level: level.effective_discount),
    "итоговая скидка": ("AS", lambda level: level.with_campaign.total_discount_rate),
}


@dataclass
class Comparison:
    """Итог сверки одной величины по всему пулу."""

    name: str
    compared: int = 0
    worst_diff: float = 0.0
    worst_contract: str | None = None

    def matches(self, tolerance: float) -> bool:
        return self.worst_diff <= tolerance


@dataclass
class VerifyReport:
    tolerance: float = DEFAULT_TOLERANCE
    contracts_compared: int = 0
    contracts_missing: list[str] = field(default_factory=list)
    period_book: int | None = None
    period_core: int | None = None
    comparisons: list[Comparison] = field(default_factory=list)
    levels_checked: int = 0
    target_equals_actual: bool = True

    @property
    def ok(self) -> bool:
        if self.contracts_missing or self.period_book != self.period_core:
            return False
        if not self.target_equals_actual:
            return False
        return all(c.matches(self.tolerance) for c in self.comparisons)

    def report(self) -> str:
        lines = [
            f"сверено договоров: {self.contracts_compared}",
            f"период выгрузки: книга {self.period_book}, ядро {self.period_core}"
            + ("" if self.period_book == self.period_core else "  РАСХОЖДЕНИЕ"),
        ]
        for c in self.comparisons:
            if c.matches(self.tolerance):
                lines.append(f"  {c.name}: совпадает ({c.compared} значений)")
            else:
                lines.append(
                    f"  {c.name}: РАСХОЖДЕНИЕ до {c.worst_diff:.4%} "
                    f"на договоре {c.worst_contract}"
                )
        if self.levels_checked:
            lines.append(f"уровней сверено: {self.levels_checked}")
            lines.append(
                "  цель равна фактической эффективной скидке"
                if self.target_equals_actual
                else "  ЦЕЛЬ НЕ РАВНА ФАКТУ — механика подбора надбавки сломана"
            )
        if self.contracts_missing:
            lines.append(
                f"  в ядре нет {len(self.contracts_missing)} договоров из книги: "
                f"{self.contracts_missing[:5]}"
            )
        lines.append("итог: сходится" if self.ok else "итог: расхождение, книгу отдавать нельзя")
        return "\n".join(lines)


def _relative(core: float, book: float) -> float:
    if book:
        return abs(core - book) / abs(book)
    return abs(core)


def verify_pool(
    book: Path,
    inputs_dir: Path,
    mapping_path: Path,
    product: str,
    tolerance: float = DEFAULT_TOLERANCE,
) -> VerifyReport:
    """Сравнить пул, собранный ядром, с колонками «Базы участников» книги."""
    if not mapping_path.exists():
        raise ModelError(
            f"нет таблицы соответствия {mapping_path}; "
            f"собери входные данные командой mk-forge prepare"
        )
    mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    to_pseudonym = {real: code for code, real in mapping.items()}

    pool: Pool = build_pool(load_inputs(inputs_dir), product)
    by_contract = {m.contract: m for m in pool.members}

    report = VerifyReport(tolerance=tolerance, period_core=pool.period_months)
    report.comparisons = [Comparison(name=name) for name in POOL_COLUMNS]

    wb = openpyxl.load_workbook(book, data_only=True)
    try:
        if CONTRACTS_SHEET not in wb.sheetnames:
            raise ModelError(f"в книге нет листа «{CONTRACTS_SHEET}»")
        ws = wb[CONTRACTS_SHEET]
        report.period_book = ws["P2"].value

        for row in range(2, LAST_ROW + 1):
            real = ws[f"A{row}"].value
            if real in (None, ""):
                continue
            pseudonym = to_pseudonym.get(str(real).strip())
            member = by_contract.get(pseudonym) if pseudonym else None
            if member is None:
                report.contracts_missing.append(str(real))
                continue
            report.contracts_compared += 1
            for comparison, (column, attribute) in zip(
                report.comparisons, POOL_COLUMNS.values(), strict=True
            ):
                value = ws[f"{column}{row}"].value
                if not isinstance(value, (int, float)):
                    continue
                comparison.compared += 1
                diff = _relative(getattr(member, attribute), value)
                if diff > comparison.worst_diff:
                    comparison.worst_diff = diff
                    comparison.worst_contract = member.contract
    finally:
        wb.close()

    return report


def verify_levels(
    book: Path,
    inputs: Inputs,
    config: CampaignConfig,
    report: VerifyReport,
) -> None:
    """Сверить раздел уровней «Расчета акции» с расчетом механики."""
    plan = build_plan(
        inputs=inputs,
        product=config.product,
        start=config.start,
        end=config.end,
        assumptions=Assumptions(
            plan_participants=config.plan_participants,
            share_without=config.share_without,
            share_with=config.share_with,
            comms_per_client=config.comms_per_client,
            new_share_without=config.new_share_without,
            large_threshold_thousand_liters=config.large_threshold_thousand_liters,
            large_shift_points=config.large_shift_points,
        ),
        targets=config.targets,
    )
    report.target_equals_actual = all(
        abs(level.effective_discount - level.target) <= report.tolerance
        for level in plan.levels
    )

    comparisons = {name: Comparison(name=name) for name in LEVEL_COLUMNS}
    wb = openpyxl.load_workbook(book, data_only=True)
    try:
        if CALCULATION_SHEET not in wb.sheetnames:
            raise ModelError(f"в книге нет листа «{CALCULATION_SHEET}»")
        ws = wb[CALCULATION_SHEET]
        for index, level in enumerate(plan.levels):
            row = LEVEL_FIRST_ROW + index
            report.levels_checked += 1
            for name, (column, extract_value) in LEVEL_COLUMNS.items():
                value = ws[f"{column}{row}"].value
                if not isinstance(value, (int, float)):
                    continue
                comparison = comparisons[name]
                comparison.compared += 1
                diff = _relative(extract_value(level), value)
                if diff > comparison.worst_diff:
                    comparison.worst_diff = diff
                    comparison.worst_contract = f"уровень {level.target * 100:g}%"
    finally:
        wb.close()

    report.comparisons.extend(comparisons.values())


def verify_all(
    book: Path,
    inputs_dir: Path,
    mapping_path: Path,
    config: CampaignConfig,
    tolerance: float = DEFAULT_TOLERANCE,
) -> VerifyReport:
    """Сверить с книгой и каждый договор, и итоги по уровням глубины скидки."""
    report = verify_pool(
        book=book,
        inputs_dir=inputs_dir,
        mapping_path=mapping_path,
        product=config.product,
        tolerance=tolerance,
    )
    verify_levels(book=book, inputs=load_inputs(inputs_dir), config=config, report=report)
    return report
