"""Сверка с эталонной книгой: расхождение должно быть найдено и названо."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from mkforge.core.models import ModelError
from mkforge.core.verify import verify_pool


def run(book: Path, anon_dir: Path, tmp_path: Path, **overrides):
    return verify_pool(
        book=book,
        inputs_dir=anon_dir,
        mapping_path=tmp_path / "mapping.json",
        product="ДТ",
        **overrides,
    )


def test_matching_book_passes(reference_book, anon_dir, tmp_path, pool):
    report = run(reference_book, anon_dir, tmp_path)
    assert report.ok
    assert report.contracts_compared == len(pool)
    assert not report.contracts_missing
    assert report.period_book == report.period_core
    assert "сходится" in report.report()


def test_changed_value_is_caught(reference_book, anon_dir, tmp_path):
    wb = openpyxl.load_workbook(reference_book)
    wb["База участников"]["O5"] = wb["База участников"]["O5"].value * 1.01
    wb.save(reference_book)

    report = run(reference_book, anon_dir, tmp_path)
    assert not report.ok
    text = report.report()
    assert "РАСХОЖДЕНИЕ" in text
    assert "среднемес. объем продукта" in text
    assert "книгу отдавать нельзя" in text


def test_deviation_within_tolerance_passes(reference_book, anon_dir, tmp_path):
    """Допуск нужен: книга хранит числа с ограниченной точностью."""
    wb = openpyxl.load_workbook(reference_book)
    ws = wb["База участников"]
    ws["K5"] = ws["K5"].value * (1 + 1e-12)
    wb.save(reference_book)

    assert run(reference_book, anon_dir, tmp_path, tolerance=1e-9).ok


def test_different_period_fails(reference_book, anon_dir, tmp_path):
    wb = openpyxl.load_workbook(reference_book)
    wb["База участников"]["P2"] = 3
    wb.save(reference_book)

    report = run(reference_book, anon_dir, tmp_path)
    assert not report.ok
    assert "РАСХОЖДЕНИЕ" in report.report()


def test_contract_absent_from_core(reference_book, anon_dir, tmp_path):
    """Договор есть в книге, но в обезличенных данных его нет."""
    path = anon_dir / "contracts.csv"
    lines = path.read_text(encoding="utf-8").splitlines()
    path.write_text("\n".join(lines[:-1]), encoding="utf-8")

    report = run(reference_book, anon_dir, tmp_path)
    assert not report.ok
    assert len(report.contracts_missing) == 1
    assert "в ядре нет" in report.report()


def test_missing_mapping_names_the_command(reference_book, anon_dir, tmp_path):
    (tmp_path / "mapping.json").unlink()
    with pytest.raises(ModelError, match="mk-forge prepare"):
        run(reference_book, anon_dir, tmp_path)


def test_book_without_pool_sheet(anon_dir, tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Не то"
    empty = tmp_path / "empty.xlsx"
    wb.save(empty)
    with pytest.raises(ModelError, match="База участников"):
        run(empty, anon_dir, tmp_path)
