"""Пересчет и проверка книги.

Полный прогон требует LibreOffice и поэтому размечен маркером: в CI без него
тест пропускается, а не падает. Все остальное проверяется без запуска.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from mkforge.config import load_config
from mkforge.core.loaders import load_inputs
from mkforge.core.validate import (
    ERROR_MARKERS,
    SOFFICE_ENV,
    RecalculationError,
    _scan_errors,
    find_soffice,
    recalc_command,
    validate,
)
from mkforge.renderers.xlsx import build


def test_profile_url_is_absolute_and_encoded(tmp_path):
    """Регрессия: от относительного URL LibreOffice молча зависает навсегда."""
    command = recalc_command(
        soffice=Path("/opt/soffice"),
        book=tmp_path / "книга.xlsx",
        out_dir=tmp_path / "пересчет",
        profile=Path("out/пересчет/профиль"),
    )
    profile_argument = next(a for a in command if a.startswith("-env:UserInstallation="))
    url = profile_argument.split("=", 1)[1]
    assert url.startswith("file:///"), "путь профиля должен быть абсолютным"
    assert "профиль" not in url, "кириллица должна быть закодирована процентами"


def test_command_passes_absolute_paths(tmp_path):
    command = recalc_command(
        soffice=Path("/opt/soffice"),
        book=Path("out/книга.xlsx"),
        out_dir=Path("out/пересчет"),
        profile=tmp_path / "профиль",
    )
    assert command[-1].startswith("/")
    assert command[command.index("--outdir") + 1].startswith("/")


def test_soffice_override_is_used(tmp_path, monkeypatch):
    fake = tmp_path / "soffice"
    fake.write_text("", encoding="utf-8")
    monkeypatch.setenv(SOFFICE_ENV, str(fake))
    assert find_soffice() == fake


def test_missing_soffice_override_is_reported(tmp_path, monkeypatch):
    monkeypatch.setenv(SOFFICE_ENV, str(tmp_path / "нет-такого"))
    with pytest.raises(RecalculationError, match=SOFFICE_ENV):
        find_soffice()


def test_error_markers_are_counted(tmp_path):
    """Ошибка в формуле не должна пройти мимо: книга с ней негодна."""
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws["A1"] = "#DIV/0!"
    ws["A2"] = "#REF!"
    ws["A3"] = "#DIV/0!"
    ws["A4"] = "обычный текст"
    assert _scan_errors(workbook) == {"#DIV/0!": 2, "#REF!": 1}


def test_all_error_markers_are_recognized():
    workbook = openpyxl.Workbook()
    ws = workbook.active
    for offset, marker in enumerate(ERROR_MARKERS, start=1):
        ws[f"A{offset}"] = marker
    assert _scan_errors(workbook) == dict.fromkeys(ERROR_MARKERS, 1)


@pytest.mark.libreoffice
def test_built_book_matches_the_core(anon_dir: Path, config_path: Path, tmp_path: Path):
    """Формулы книги и расчетное ядро должны дать одинаковые числа."""
    pytest.importorskip("openpyxl")
    try:
        find_soffice()
    except RecalculationError:
        pytest.skip("LibreOffice не установлен")

    book = tmp_path / "книга.xlsx"
    inputs = load_inputs(anon_dir)
    config = load_config(config_path)
    build(inputs=inputs, config=config, path=book)

    report = validate(
        book=book, inputs=inputs, config=config, work_dir=tmp_path / "пересчет"
    )
    assert not report.formula_errors, report.report()
    assert report.target_equals_actual, report.report()
    assert report.ok, report.report()
    assert report.levels_checked == len(config.targets)
    assert all(c.compared > 0 for c in report.comparisons)


@pytest.mark.libreoffice
def test_distributed_book_matches_the_core(
    anon_dir: Path, distributed_config_path: Path, tmp_path: Path
):
    """Третий путь: формулы раздела распределения против расчетного ядра.

    Проверяется и порог безубыточности: книга берет его формулой, ядро — решением
    того же уравнения, и расходиться им нельзя.
    """
    try:
        find_soffice()
    except RecalculationError:
        pytest.skip("LibreOffice не установлен")

    config = load_config(distributed_config_path)
    inputs = load_inputs(anon_dir)
    book = tmp_path / "книга-распределение.xlsx"
    build(inputs=inputs, config=config, path=book)

    report = validate(
        book=book, inputs=inputs, config=config, work_dir=tmp_path / "пересчет"
    )
    assert report.ok, report.report()
    assert report.depths_checked == len(config.distribution.shares) + 1
    assert report.depth_equals_actual
    assert "совпадает с ядром" in report.breakeven
