"""Витрина, расчет и справочники: адреса ввода, уровни, печать допущений."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from mkforge.config import load_config
from mkforge.core.loaders import load_inputs
from mkforge.renderers import calculation as calc
from mkforge.renderers import dictionaries, showcase
from mkforge.renderers.xlsx import build


@pytest.fixture
def config(config_path: Path):
    return load_config(config_path)


@pytest.fixture
def sheets(anon_dir: Path, config, tmp_path: Path):
    path = tmp_path / "книга.xlsx"
    build(inputs=load_inputs(anon_dir), config=config, path=path)
    workbook = openpyxl.load_workbook(path)
    yield workbook
    workbook.close()


def test_input_cells_are_at_reference_addresses(sheets, config):
    """На эти ячейки ссылается весь расчет, поэтому адреса менять нельзя."""
    ws = sheets[showcase.SHEET]
    assert ws[f"B{showcase.START}"].value.date() == config.start
    assert ws[f"B{showcase.END}"].value.date() == config.end
    assert ws[f"B{showcase.SHARE_WITHOUT}"].value == config.share_without
    assert ws[f"B{showcase.SHARE_WITH}"].value == config.share_with
    assert ws[f"B{showcase.COMMS}"].value == config.comms_per_client
    assert ws[f"B{showcase.SELECTED_LEVEL}"].value == config.showcase_level


def test_selected_level_is_constrained(sheets, config):
    """Уровень выбирается из набора, а не вводится произвольно."""
    ws = sheets[showcase.SHEET]
    ranges = [str(v.sqref) for v in ws.data_validations.dataValidation]
    assert f"B{showcase.SELECTED_LEVEL}" in ranges


def test_level_table_has_a_row_per_target(sheets, config):
    ws = sheets[calc.SHEET]
    for index, target in enumerate(config.targets):
        row = calc.LEVELS_FIRST_ROW + index
        assert ws[f"AQ{row}"].value == pytest.approx(target)
        assert ws[f"A{row}"].value == f"Уровень {target * 100:g}%"
    beyond = calc.LEVELS_FIRST_ROW + len(config.targets)
    assert ws[f"AQ{beyond}"].value is None


def test_targets_are_stored_as_fractions(sheets, config):
    ws = sheets[calc.SHEET]
    for index in range(len(config.targets)):
        assert ws[f"AQ{calc.LEVELS_FIRST_ROW + index}"].value < 1


def test_markup_formula_solves_for_the_target(sheets):
    """Надбавка выводится из цели, а не задается руками."""
    formula = sheets[calc.SHEET][f"B{calc.LEVELS_FIRST_ROW}"].value
    row = calc.LEVELS_FIRST_ROW
    assert formula == (
        f"=IF(T{row}+AV{row}>0,(AQ{row}*T{row}-AT{row}+AU{row})/(T{row}+AV{row}),0)"
    )


def test_invariant_columns_are_present(sheets, config):
    """Цель и факт стоят рядом: расхождение между ними видно глазами."""
    ws = sheets[calc.SHEET]
    assert ws[f"AQ{calc.LEVELS_HEADER_ROW}"].value.startswith("Целевая")
    assert ws[f"AR{calc.LEVELS_HEADER_ROW}"].value.startswith("Эффективная скидка факт")
    note = ws[f"A{calc.levels_last_row(len(config.targets)) + 1}"].value
    assert "AQ и AR" in note


def test_indicators_read_the_selected_level(sheets):
    """Показатели наверху листа берутся из выбранного уровня, а не считаются заново."""
    ws = sheets[calc.SHEET]
    for address in ("B6", "B7", "B9", f"B{calc.MARKUP_CURRENT}"):
        formula = ws[address].value
        assert "INDEX(" in formula and "MATCH(" in formula
        assert "$AQ$" in formula


def test_new_clients_never_go_negative(sheets):
    """План ниже текущего пула не должен давать отрицательных новых."""
    formula = sheets[calc.SHEET][f"B{calc.NEW_COUNT}"].value
    assert formula.startswith("=MAX(0,")


def test_rules_are_printed_with_substitutions(sheets, config):
    ws = sheets[showcase.SHEET]
    first = ws[f"B{showcase.RULES_HEADER_ROW + 1}"].value
    assert first == config.rules_text()[0]
    assert "{" not in first


def test_assumptions_are_printed(sheets, config):
    """Обещание выполнено: экономист видит допущения, не читая формулы."""
    ws = sheets[showcase.SHEET]
    printed = "\n".join(
        str(cell.value) for row in ws.iter_rows(min_col=1, max_col=3) for cell in row if cell.value
    )
    for assumption in config.assumptions:
        assert assumption.name.replace("_", " ") in printed
        assert assumption.why.strip() in printed


def test_level_summary_points_at_the_calculation(sheets, config):
    ws = sheets[showcase.SHEET]
    for index in range(len(config.targets)):
        row = showcase.LEVELS_FIRST_ROW + index
        source = calc.LEVELS_FIRST_ROW + index
        assert ws[f"B{row}"].value == f"='{calc.SHEET}'!$AQ${source}"


def test_results_show_both_scenarios(sheets):
    ws = sheets[showcase.SHEET]
    for _, row, _, _, _ in showcase.RESULTS:
        assert ws[f"B{row}"].value.startswith("=IFERROR(INDEX(")
        assert ws[f"C{row}"].value.startswith("=IFERROR(INDEX(")
        assert ws[f"D{row}"].value is not None


def test_campaign_card_describes_the_mechanic(sheets, config):
    ws = sheets[dictionaries.CAMPAIGN_SHEET]
    values = [str(ws[f"B{row}"].value) for row in range(3, 13)]
    assert config.name in values
    assert config.mechanic in values
    assert config.product in values


def test_reference_lists_products_from_the_export(sheets, anon_dir):
    ws = sheets[dictionaries.REFERENCE_SHEET]
    inputs = load_inputs(anon_dir)
    expected = sorted({(t.product, t.product_class) for t in inputs.transactions if t.is_classified})
    written = []
    for offset in range(len(expected)):
        row = dictionaries.PRODUCTS_FIRST_ROW + offset
        written.append((ws[f"A{row}"].value, ws[f"B{row}"].value))
    assert written == expected


def test_parameter_blocks_do_not_overwrite_each_other(sheets, anon_dir: Path):
    """Блоки листа параметров стоят рядом: ни один не должен затирать другой.

    Ввод блока «Состав новых» когда-то стоял в колонке L и затирал пять строк
    прогноза маржи. Проявилось бы это только при другом периоде акции, поэтому
    проверяем прямо: все строки прогноза на месте и совпадают со входом.
    """
    from mkforge.renderers import parameters as params

    inputs = load_inputs(anon_dir)
    ws = sheets[params.SHEET]
    first = params.FORECAST_FIRST_ROW
    for offset, line in enumerate(inputs.margin):
        row = first + offset
        assert ws[f"L{row}"].value == line.month, f"месяц в строке {row} затерт"
        assert ws[f"M{row}"].value == line.product, f"продукт в строке {row} затерт"
        assert ws[f"N{row}"].value == line.margin, f"маржа в строке {row} затерта"
        assert ws[f"O{row}"].value == line.region, f"регион в строке {row} затерт"


@pytest.fixture
def distributed(anon_dir: Path, distributed_config_path: Path, tmp_path: Path):
    """Книга, собранная с распределением глубины скидки."""
    config = load_config(distributed_config_path)
    path = tmp_path / "книга-распределение.xlsx"
    build(inputs=load_inputs(anon_dir), config=config, path=path)
    workbook = openpyxl.load_workbook(path)
    yield workbook, config
    workbook.close()


def test_forecast_sheet_appears_with_a_distribution(distributed, anon_dir: Path):
    from mkforge.campaigns.motivational import forecast_pool_from
    from mkforge.renderers import forecast as fc

    workbook, config = distributed
    assert fc.SHEET in workbook.sheetnames
    ws = workbook[fc.SHEET]

    expected = forecast_pool_from(load_inputs(anon_dir), config)
    rows = 0
    row = fc.FIRST_ROW
    while str(ws[f"{fc.NUMBER}{row}"].value or "").startswith(fc.NUMBER_PREFIX):
        rows += 1
        row += 1
    assert rows == expected.size, "строк должно быть столько же, сколько в прогнозе ядра"
    assert isinstance(ws[f"{fc.TERM}{fc.FIRST_ROW}"].value, str), "срок должен быть формулой"
    assert ws[f"{fc.TONS}{fc.FIRST_ROW}"].value.startswith("="), "объем — формула по сегменту"


def test_no_forecast_sheet_without_a_distribution(sheets):
    from mkforge.renderers import forecast as fc

    assert fc.SHEET not in sheets.sheetnames


def test_distribution_shares_are_input_cells(distributed):
    """Доли — это ввод: экономист вправе поменять их в книге, не пересобирая."""
    from mkforge.renderers import distribution as dist

    workbook, config = distributed
    ws = workbook["Расчет акции"]
    ordered = config.distribution.deepest_first()
    total = 0.0
    for offset, share in enumerate(ordered.shares):
        row = dist.FIRST_ROW + offset
        assert ws[f"{dist.DEPTH}{row}"].value == pytest.approx(share.depth)
        cell = ws[f"{dist.SHARE}{row}"]
        assert cell.value == pytest.approx(share.share)
        assert cell.style == "mk_input", "доля должна быть помечена как ввод"
        total += cell.value
    assert total == pytest.approx(1.0)


def test_handout_mode_is_an_input_cell(distributed):
    from mkforge.renderers import distribution as dist

    workbook, config = distributed
    ws = workbook["Расчет акции"]
    assert ws[dist.MODE].value == (1 if config.distribution.by_volume else 0)
    assert ws[dist.MODE].style == "mk_input_count"


def test_depth_is_never_a_hardcoded_number_per_contract(distributed):
    """Глубина договора считается в книге, а не проставлена сборщиком."""
    from mkforge.renderers import participants as pool

    workbook, _ = distributed
    ws = workbook[pool.SHEET]
    for row in (pool.FIRST_ROW, pool.FIRST_ROW + 1):
        for column in (pool.RANK, pool.CUMULATIVE, pool.DEPTH, pool.MARKUP):
            value = ws[f"{column}{row}"].value
            assert isinstance(value, str) and value.startswith("="), (
                f"{column}{row} должна быть формулой, а не значением"
            )


def test_showcase_shows_the_total_and_the_threshold(distributed):
    """Порог важнее результата: результат зависит от допущения, порог — нет."""
    from mkforge.renderers import distribution as dist
    from mkforge.renderers import showcase

    workbook, config = distributed
    ws = workbook[showcase.SHEET]
    count = len(config.distribution.shares)
    total = dist.total_row(count)
    breakeven = list(dist.breakeven_rows(count).values())

    panel = [
        ws[f"{showcase.SUMMARY_VALUE}{showcase.SUMMARY_FIRST_ROW + offset}"].value
        for offset in range(8)
    ]
    assert any(f"${dist.PAYBACK}${total}" in str(value) for value in panel)
    assert any(f"$B${breakeven[4]}" in str(value) for value in panel)
    assert any(f"$B${breakeven[5]}" in str(value) for value in panel)


def test_summary_sheet_comes_first_and_says_the_verdict(distributed):
    """Первое, что видит человек: ответ, а не расчет.

    Претензия была прямая — «очень мудрит, нужно чтобы понятно было». Числа
    от этого не изменились, но лист итога обязан быть первым и обязан
    отвечать словами, а не только цифрами.
    """
    from mkforge.renderers import distribution as dist
    from mkforge.renderers import summary

    workbook, config = distributed
    assert workbook.sheetnames[0] == summary.SHEET
    ws = workbook[summary.SHEET]

    texts = [
        str(ws.cell(row=row, column=column).value or "")
        for row in range(1, 30)
        for column in (1, 2)
    ]
    joined = " ".join(texts)
    assert "минус" in joined, "вердикт про минус должен быть словами"
    total = dist.total_row(len(config.distribution.shares))
    assert any(f"${dist.PAYBACK}${total}" in text for text in texts), (
        "окупаемость должна ссылаться на итог раздела распределения, "
        "а не считаться на листе заново"
    )


def test_service_columns_are_hidden_not_deleted(distributed):
    """Спрятанное можно раскрыть, удаленное пришлось бы пересобирать."""
    from mkforge.renderers import forecast as fc
    from mkforge.renderers import participants as pool

    workbook, _ = distributed
    forecast_sheet = workbook[fc.SHEET]
    for column in (fc.TON_MONTHS, fc.RANK, fc.CUMULATIVE):
        assert forecast_sheet.column_dimensions[column].hidden, column
    assert not forecast_sheet.column_dimensions[fc.DEPTH].hidden, (
        "глубина скидки — то, зачем лист и нужен"
    )

    pool_sheet = workbook[pool.SHEET]
    assert pool_sheet.column_dimensions[pool.RANK].hidden
    assert pool_sheet.column_dimensions[pool.CUMULATIVE].hidden
    assert not pool_sheet.column_dimensions[pool.DEPTH].hidden


def test_reference_level_table_is_hidden_with_a_distribution(distributed):
    """Против каждой скидки в ней стоит весь пул, и ее принимают за сценарий акции.

    Претензия была конкретная: «по каждой скидке участвует весь план, тогда как план
    должен распределиться». Таблица верна для своей задачи, но лежит первой
    на пути, и подписи «справочно» оказалось недостаточно.
    """
    workbook, config = distributed
    ws = workbook[calc.SHEET]
    last = calc.levels_last_row(len(config.targets))
    for row in range(calc.LEVELS_SECTION_ROW, last + 1):
        assert ws.row_dimensions[row].hidden, f"строка {row} должна быть скрыта"
    pointer = ws[f"A{calc.LEVELS_SECTION_ROW - 1}"].value
    assert "раздел 6" in str(pointer), "на месте таблицы должно остаться указание"


def test_reference_level_table_stays_visible_without_a_distribution(sheets, config):
    """В знакомой книге эта таблица и есть сценарий: прятать ее нельзя."""
    ws = sheets[calc.SHEET]
    last = calc.levels_last_row(len(config.targets))
    for row in range(calc.LEVELS_SECTION_ROW, last + 1):
        assert not ws.row_dimensions[row].hidden, f"строка {row} скрыта напрасно"
