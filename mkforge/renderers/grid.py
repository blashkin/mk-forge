"""Мелкие помощники для раскладки листов.

Адреса ячеек повторяют эталонную книгу: так собранную книгу можно сверить
с эталоном механически, ячейка к ячейке, а не сопоставляя величины руками.
"""

from __future__ import annotations

from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

SECTION_HEIGHT = 20


def put(
    ws: Worksheet,
    address: str,
    value: Any = None,
    style: str | None = None,
    number_format: str | None = None,
) -> None:
    """Записать значение или формулу в ячейку, задав стиль по имени."""
    cell = ws[address]
    if value is not None:
        cell.value = value
    if style:
        cell.style = style
    if number_format:
        cell.number_format = number_format


def section(ws: Worksheet, row: int, title: str, span: int, first_column: int = 1) -> None:
    """Заголовок раздела: плашка через несколько колонок.

    `first_column` нужен потому, что на листе параметров блоки стоят рядом
    по горизонтали и делят одну строку.
    """
    for offset in range(span):
        cell = ws.cell(row=row, column=first_column + offset)
        cell.style = "mk_section"
        if offset == 0:
            cell.value = title
    ws.row_dimensions[row].height = SECTION_HEIGHT


def headers(ws: Worksheet, row: int, labels: dict[str, str]) -> None:
    """Шапка таблицы: колонка -> подпись."""
    for column, label in labels.items():
        put(ws, f"{column}{row}", label, "mk_header")


def widths(ws: Worksheet, mapping: dict[str, int]) -> None:
    for column, width in mapping.items():
        ws.column_dimensions[column].width = width


def column_range(sheet: str, column: str, last_row: int, first_row: int = 2) -> str:
    """Абсолютная ссылка на колонку другого листа от первой строки данных до последней.

    Верхняя граница считается от фактического числа строк. В эталонной книге она
    была вписана руками и разъехалась между колонками — повторять это нельзя.
    """
    return f"'{sheet}'!${column}${first_row}:${column}${last_row}"


def letter(index: int) -> str:
    return get_column_letter(index)
