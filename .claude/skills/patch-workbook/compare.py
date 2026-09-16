"""Сличение чужой книги с ее правкой: что изменилось, кроме расчета.

Печатает адреса, а не содержимое. Значения ячеек сравниваются по хешу и
наружу не выводятся: отчет можно показывать и пересылать, данных клиентов
в нем нет. Формулы показываются только по флагу --формулы.

Код возврата: 0, если форма книги не тронута (листы, размеры, именованные
диапазоны, проверки ввода, объединения, автофильтр, закрепление, скрытые
строки и колонки) — расхождения в формулах и значениях при этом допустимы
и перечисляются. 1 — если изменилась форма.
"""

from __future__ import annotations

import argparse
import hashlib
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl


def _digest(value: object) -> str:
    """Отпечаток значения: сравнение точное, само значение не сохраняется."""
    return hashlib.sha256(repr(value).encode("utf-8")).hexdigest()[:16]


@dataclass(frozen=True)
class Shape:
    """Форма листа: все, что не является расчетом."""

    state: str
    max_row: int
    max_column: int
    freeze: str | None
    autofilter: str | None
    merged: tuple[str, ...]
    validations: tuple[str, ...]
    hidden_columns: tuple[str, ...]
    hidden_rows: tuple[int, ...]

    def differences(self, other: "Shape") -> list[str]:
        lines = []
        for field_name, label in (
            ("state", "видимость листа"),
            ("max_row", "число строк"),
            ("max_column", "число колонок"),
            ("freeze", "закрепление областей"),
            ("autofilter", "автофильтр"),
            ("merged", "объединенные ячейки"),
            ("validations", "проверки ввода"),
            ("hidden_columns", "скрытые колонки"),
            ("hidden_rows", "скрытые строки"),
        ):
            was, now = getattr(self, field_name), getattr(other, field_name)
            if was != now:
                if isinstance(was, tuple):
                    lines.append(
                        f"    {label}: было {len(was)}, стало {len(now)}"
                        f" — добавлено {sorted(set(now) - set(was))[:10]}"
                        f", убрано {sorted(set(was) - set(now))[:10]}"
                    )
                else:
                    lines.append(f"    {label}: было {was!r}, стало {now!r}")
        return lines


@dataclass(frozen=True)
class Sheet:
    shape: Shape
    formulas: dict[str, str]
    values: dict[str, str]


@dataclass
class Book:
    path: Path
    sheets: dict[str, Sheet] = field(default_factory=dict)
    names: dict[str, str] = field(default_factory=dict)


def fingerprint(path: Path) -> Book:
    """Снять отпечаток книги: форма листов, формулы, хеши значений."""
    workbook = openpyxl.load_workbook(path)
    try:
        book = Book(path=path)
        for name, definition in workbook.defined_names.items():
            book.names[name] = str(definition.value)

        for ws in workbook.worksheets:
            formulas: dict[str, str] = {}
            values: dict[str, str] = {}
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas[cell.coordinate] = cell.value
                    else:
                        values[cell.coordinate] = _digest(cell.value)

            shape = Shape(
                state=ws.sheet_state,
                max_row=ws.max_row,
                max_column=ws.max_column,
                freeze=ws.freeze_panes,
                autofilter=ws.auto_filter.ref,
                merged=tuple(sorted(str(area) for area in ws.merged_cells.ranges)),
                validations=tuple(
                    sorted(
                        f"{rule.type}:{rule.sqref}"
                        for rule in ws.data_validations.dataValidation
                    )
                ),
                hidden_columns=tuple(
                    sorted(key for key, dim in ws.column_dimensions.items() if dim.hidden)
                ),
                hidden_rows=tuple(
                    sorted(key for key, dim in ws.row_dimensions.items() if dim.hidden)
                ),
            )
            book.sheets[ws.title] = Sheet(shape=shape, formulas=formulas, values=values)
        return book
    finally:
        workbook.close()


def _changes(was: dict[str, str], now: dict[str, str]) -> tuple[list[str], list[str], list[str]]:
    added = sorted(set(now) - set(was))
    removed = sorted(set(was) - set(now))
    changed = sorted(key for key in set(was) & set(now) if was[key] != now[key])
    return added, removed, changed


def compare(source: Book, patched: Book, show_formulas: bool, limit: int) -> tuple[str, bool]:
    """Отчет о различиях и признак того, что форма книги не тронута."""
    lines: list[str] = []
    form_intact = True

    lost = sorted(set(source.sheets) - set(patched.sheets))
    gained = sorted(set(patched.sheets) - set(source.sheets))
    if lost or gained:
        form_intact = False
        lines.append(f"Листы: убрано {lost}, добавлено {gained}")
    if list(source.sheets) != list(patched.sheets):
        form_intact = False
        lines.append("Порядок листов изменился")

    if source.names != patched.names:
        form_intact = False
        lines.append(
            f"Именованные диапазоны: было {len(source.names)}, стало {len(patched.names)}"
        )

    for title in source.sheets:
        if title not in patched.sheets:
            continue
        was, now = source.sheets[title], patched.sheets[title]
        block: list[str] = []

        shape_lines = was.shape.differences(now.shape)
        if shape_lines:
            form_intact = False
            block.append("  форма:")
            block.extend(shape_lines)

        added, removed, changed = _changes(was.formulas, now.formulas)
        if added or removed or changed:
            block.append(
                f"  формулы: добавлено {len(added)}, убрано {len(removed)},"
                f" изменено {len(changed)}"
            )
            for label, cells in (("добавлены", added), ("убраны", removed), ("изменены", changed)):
                if not cells:
                    continue
                block.append(f"    {label}: {', '.join(cells[:limit])}"
                             + (" ..." if len(cells) > limit else ""))
            if show_formulas:
                for cell in changed[:limit]:
                    block.append(f"    {cell}: было {was.formulas[cell]}")
                    block.append(f"    {cell}: стало {now.formulas[cell]}")
                for cell in added[:limit]:
                    block.append(f"    {cell}: {now.formulas[cell]}")

        added, removed, changed = _changes(was.values, now.values)
        if added or removed or changed:
            block.append(
                f"  значения: добавлено {len(added)}, убрано {len(removed)},"
                f" изменено {len(changed)}"
            )
            for label, cells in (("добавлены", added), ("убраны", removed), ("изменены", changed)):
                if not cells:
                    continue
                block.append(f"    {label}: {', '.join(cells[:limit])}"
                             + (" ..." if len(cells) > limit else ""))

        if block:
            lines.append(f"Лист «{title}»:")
            lines.extend(block)

    if not lines:
        lines.append("Различий нет: книги совпадают по форме, формулам и значениям.")
    return "\n".join(lines), form_intact


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Сличить исходную книгу с правкой: что изменилось, кроме расчета"
    )
    parser.add_argument("исходник", type=Path, help="книга, которую прислали")
    parser.add_argument("правка", type=Path, help="книга после правки")
    parser.add_argument("--формулы", action="store_true", dest="formulas",
                        help="показать текст формул, а не только адреса")
    parser.add_argument("--предел", type=int, default=20, dest="limit",
                        help="сколько адресов печатать на список")
    args = parser.parse_args(argv)

    source = fingerprint(args.исходник)
    patched = fingerprint(args.правка)
    report, form_intact = compare(source, patched, args.formulas, args.limit)
    print(report)
    print()
    print("Форма книги не тронута." if form_intact
          else "ФОРМА КНИГИ ИЗМЕНИЛАСЬ — это запрещено правилом поставки.")
    return 0 if form_intact else 1


if __name__ == "__main__":
    raise SystemExit(main())
