"""Стили книги, описанные кодом.

Эталонная книга копировала оформление из исходного файла, поэтому оно нигде
не было записано. Здесь стили заданы явно: книга собирается без шаблона,
и в CI ей ничего не нужно, кроме кода.

Синие цифры на желтом фоне — то, что человек вправе менять руками.
Все остальное считается формулами, и трогать это незачем.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side

FONT = "Calibri"

SECTION_FILL = "FF1F4E79"
HEADER_FILL = "FFDCE6F1"
INPUT_FILL = "FFFFF2CC"
NOTE_COLOR = "FF7F7F7F"
INPUT_COLOR = "FF0000FF"

# Форматы чисел. Проценты хранятся долями, поэтому формат делает из 0.025 «2,50%».
PERCENT = "0.00%"
PERCENT_FINE = "0.000%"
MONEY = "#,##0"
TONS = "#,##0.000"
COUNT = "#,##0"
RATIO = "#,##0.00"
DATE = "DD.MM.YYYY"

_thin = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def _style(name: str, **kwargs) -> NamedStyle:
    return NamedStyle(name=name, **kwargs)


def named_styles() -> list[NamedStyle]:
    """Стили книги. Регистрируются один раз на книгу, дальше по имени."""
    return [
        _style(
            "mk_section",
            font=Font(name=FONT, size=12, bold=True, color="FFFFFFFF"),
            fill=PatternFill("solid", fgColor=SECTION_FILL),
            alignment=Alignment(horizontal="left", vertical="center"),
        ),
        _style(
            "mk_header",
            font=Font(name=FONT, size=11, bold=True),
            fill=PatternFill("solid", fgColor=HEADER_FILL),
            alignment=CENTER,
            border=BORDER,
        ),
        _style(
            "mk_label",
            font=Font(name=FONT, size=11, bold=True),
            alignment=WRAP,
            border=BORDER,
        ),
        _style("mk_text", font=Font(name=FONT, size=11), alignment=WRAP, border=BORDER),
        _style(
            "mk_note",
            font=Font(name=FONT, size=10, italic=True, color=NOTE_COLOR),
            alignment=WRAP,
        ),
        _style(
            "mk_input",
            font=Font(name=FONT, size=11, bold=True, color=INPUT_COLOR),
            fill=PatternFill("solid", fgColor=INPUT_FILL),
            alignment=RIGHT,
            border=BORDER,
            number_format=PERCENT,
        ),
        _style(
            "mk_input_date",
            font=Font(name=FONT, size=11, bold=True, color=INPUT_COLOR),
            fill=PatternFill("solid", fgColor=INPUT_FILL),
            alignment=RIGHT,
            border=BORDER,
            number_format=DATE,
        ),
        _style(
            "mk_input_count",
            font=Font(name=FONT, size=11, bold=True, color=INPUT_COLOR),
            fill=PatternFill("solid", fgColor=INPUT_FILL),
            alignment=RIGHT,
            border=BORDER,
            number_format=COUNT,
        ),
        _style("mk_pct", font=Font(name=FONT, size=11), alignment=RIGHT, border=BORDER, number_format=PERCENT),
        _style("mk_pct_fine", font=Font(name=FONT, size=11), alignment=RIGHT, border=BORDER, number_format=PERCENT_FINE),
        _style("mk_money", font=Font(name=FONT, size=11), alignment=RIGHT, border=BORDER, number_format=MONEY),
        _style("mk_tons", font=Font(name=FONT, size=11), alignment=RIGHT, border=BORDER, number_format=TONS),
        _style("mk_count", font=Font(name=FONT, size=11), alignment=RIGHT, border=BORDER, number_format=COUNT),
        _style("mk_ratio", font=Font(name=FONT, size=11), alignment=RIGHT, border=BORDER, number_format=RATIO),
        _style("mk_date", font=Font(name=FONT, size=11), alignment=RIGHT, border=BORDER, number_format=DATE),
    ]


def register(workbook) -> None:
    """Добавить стили в книгу. Повторная регистрация игнорируется."""
    existing = set(workbook.named_styles)
    for style in named_styles():
        if style.name not in existing:
            workbook.add_named_style(style)
